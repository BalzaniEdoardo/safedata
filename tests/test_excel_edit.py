import datetime as dt

from conftest import make_row
from openpyxl import load_workbook

from safedata.changeset import apply
from safedata.excel_edit import changeset_from_edits, export, read_edits


def edit_workbook(path, fn):
    wb = load_workbook(path)
    ws = wb["Records"]
    header = [c.value for c in ws[1]]
    fn(ws, {name: i + 1 for i, name in enumerate(header)})
    wb.save(path)


def test_excel_roundtrip(tmp_path, schema):
    base = [make_row(i) for i in range(1, 6)]
    path = tmp_path / "edit.xlsx"
    export(base, schema, "abc123", path)

    # Untouched workbook -> no changes
    sheet = read_edits(path, schema)
    assert sheet.base_commit == "abc123"
    cs, errors = changeset_from_edits(base, sheet, schema)
    assert errors == [] and cs.is_empty()

    def edits(ws, col):
        ws.cell(row=2, column=col["value"], value=42)                   # record 1: change value
        ws.cell(row=3, column=col["_action"], value="delete")           # record 2: delete
        ws.delete_rows(4)                                               # record 3: row removed -> kept
        new = ws.max_row + 1                                            # new record, Excel-typed values
        ws.cell(row=new, column=col["date"], value=dt.datetime(2021, 5, 6))
        ws.cell(row=new, column=col["site"], value="south")
        ws.cell(row=new, column=col["category"], value="sample")
        ws.cell(row=new, column=col["observer"], value="Luca")
        ws.cell(row=new, column=col["verified"], value=False)

    edit_workbook(path, edits)
    cs, errors = changeset_from_edits(base, read_edits(path, schema), schema)
    assert errors == []
    assert cs.updates == {"1": {"value": ("1.5", "42")}}
    assert list(cs.deletes) == ["2"]
    assert cs.adds == [{"id": "", "date": "2021-05-06", "site": "South", "category": "sample",
                        "observer": "Luca", "value": "", "verified": "false", "notes": ""}]

    result = apply(base, cs, schema)
    assert not result.conflicts
    assert [r["id"] for r in result.rows] == ["1", "3", "4", "5", "6"]


def test_excel_rejects_unknown_and_duplicate_ids(tmp_path, schema):
    base = [make_row(1), make_row(2)]
    path = tmp_path / "edit.xlsx"
    export(base, schema, "abc123", path)

    def edits(ws, col):
        ws.cell(row=3, column=col["id"], value=1)       # duplicate of row 2
        ws.cell(row=4, column=col["id"], value=99)      # unknown id
        ws.cell(row=4, column=col["observer"], value="X")
        ws.cell(row=5, column=col["_action"], value="remove")
        ws.cell(row=5, column=col["observer"], value="Y")

    edit_workbook(path, edits)
    _, errors = changeset_from_edits(base, read_edits(path, schema), schema)
    text = "\n".join(errors)
    assert "record 1 also appears on row 2" in text
    assert "unknown _action" in text
    assert "Record 99 does not exist" in text
