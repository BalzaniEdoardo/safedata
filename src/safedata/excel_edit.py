"""Excel round-trip editing: export an editable workbook, read edits back as a changeset."""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .changeset import Changeset, diff
from .schema import Column, Schema, canon
from .store import Rows, canonical_rows, index

FORMAT_TAG = "safedata-edit-1"
ACTION = "_action"
DATA_SHEET = "Records"
SPARE_ROWS = 1000  # rows below the data that still get dropdowns and checks

HELP_TEXT = [
    "How to edit the records",
    "",
    "Change a value: just type over it in the Records sheet.",
    "Add a record: fill in an empty row at the bottom and leave the id blank. It gets a number when you submit.",
    "Delete a record: choose 'delete' in the _action column of that row. Removing the row does NOT delete it.",
    "Sorting and filtering this sheet is fine; nothing is lost.",
    "Do not change the id column.",
    "",
    "When you are done: save and close this file, then double-click 'submit'.",
    "You will see a summary of your changes and be asked to confirm.",
]


@dataclass
class EditSheet:
    base_commit: str
    rows: Rows = field(default_factory=list)
    delete_keys: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def export(rows: Rows, schema: Schema, base_commit: str, path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    headers = [ACTION] + schema.names
    ws.append(headers)
    for row in canonical_rows(rows, schema):
        ws.append([""] + [_to_excel(row[c.name], c) for c in schema.columns])

    last_row = len(rows) + 1 + SPARE_ROWS
    lists = wb.create_sheet("_lists")
    lists.sheet_state = "hidden"
    n_lists = 0

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    locked_fill = PatternFill("solid", fgColor="E7E6E6")
    for i, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = header_font, header_fill
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, len(name) + 4)
        target = f"{letter}2:{letter}{last_row}"

        if name == ACTION:
            cell.comment = Comment("Choose 'delete' to remove this record.", "safedata")
            _add_validation(ws, DataValidation(type="list", formula1='"delete"', allow_blank=True), target)
            continue
        col = schema.column(name)
        if col.description:
            cell.comment = Comment(col.description, "safedata")
        if name == schema.key:
            cell.comment = Comment("Filled in automatically. Leave blank for new records; never change it.", "safedata")
            for r in range(2, len(rows) + 2):
                ws.cell(row=r, column=i).fill = locked_fill
        if col.type == "choice":
            n_lists += 1
            list_col = get_column_letter(n_lists)
            for j, choice in enumerate(col.choices, start=1):
                lists[f"{list_col}{j}"] = choice
            formula = f"'_lists'!${list_col}$1:${list_col}${len(col.choices)}"
            _add_validation(ws, _strict(DataValidation(type="list", formula1=formula, allow_blank=True),
                                        f"Pick one of: {', '.join(col.choices)}"), target)
        elif col.type == "date":
            ws.column_dimensions[letter].width = 13
            for r in range(2, last_row + 1):
                ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
            low = dt.date.fromisoformat(col.min) if col.min else dt.date(1900, 1, 2)
            high = dt.date.fromisoformat(col.max) if col.max else dt.date(9999, 12, 31)
            dv = DataValidation(type="date", operator="between", allow_blank=True,
                                formula1=f"DATE({low.year},{low.month},{low.day})",
                                formula2=f"DATE({high.year},{high.month},{high.day})")
            _add_validation(ws, _strict(dv, "Enter a date, e.g. 2024-03-15."), target)
        elif col.type in ("int", "float"):
            dv = DataValidation(type="whole" if col.type == "int" else "decimal", operator="between",
                                allow_blank=True, formula1=col.min or "-1E+300", formula2=col.max or "1E+300")
            _add_validation(ws, _strict(dv, f"Enter a {'whole ' if col.type == 'int' else ''}number"
                                            + (f" of at least {col.min}" if col.min else "") + "."), target)
        elif col.type == "string":
            ws.column_dimensions[letter].width = 30 if name == "notes" else 18

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    help_ws = wb.create_sheet("Help")
    for line in HELP_TEXT:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True, size=14)
    help_ws.column_dimensions["A"].width = 110

    info = wb.create_sheet("_info")
    info.sheet_state = "veryHidden"
    info.append(["format", FORMAT_TAG])
    info.append(["base_commit", base_commit])

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _strict(dv: DataValidation, message: str) -> DataValidation:
    dv.showErrorMessage = True
    dv.errorTitle = "Invalid value"
    dv.error = message
    return dv


def _add_validation(ws, dv: DataValidation, target: str) -> None:
    ws.add_data_validation(dv)
    dv.add(target)


def _to_excel(value: str, col: Column):
    if value == "":
        return None
    if col.type == "int":
        return int(value)
    if col.type == "float":
        return float(value)
    if col.type == "date":
        return dt.date.fromisoformat(value)
    if col.type == "bool":
        return value == "true"
    return value


def read_edits(path: str | Path, schema: Schema) -> EditSheet:
    wb = load_workbook(path, data_only=True)
    if "_info" not in wb.sheetnames or DATA_SHEET not in wb.sheetnames:
        raise ValueError(f"{path} is not an editing workbook made by 'safedata edit'")
    info = {r[0]: r[1] for r in wb["_info"].iter_rows(values_only=True) if r and r[0]}
    if info.get("format") != FORMAT_TAG:
        raise ValueError(f"{path} was made by an incompatible version of the tools")
    sheet = EditSheet(base_commit=str(info["base_commit"]))

    ws = wb[DATA_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [n for n in [ACTION] + schema.names if n not in header]
    if missing:
        sheet.errors.append(f"The Records sheet is missing columns: {', '.join(missing)}")
        return sheet
    pos = {name: header.index(name) for name in [ACTION] + schema.names}

    seen: dict[str, int] = {}
    for line, values in enumerate(rows, start=2):
        values = list(values) + [None] * (len(header) - len(values))
        cells = {name: values[i] for name, i in pos.items()}
        if all(v is None or str(v).strip() == "" for v in cells.values()):
            continue
        action = str(cells[ACTION] or "").strip().lower()
        row = {c.name: canon(cells[c.name], c) for c in schema.columns}
        key = row[schema.key]
        if key:
            if key in seen:
                sheet.errors.append(f"Row {line}: record {key} also appears on row {seen[key]}")
                continue
            seen[key] = line
        if action == "delete":
            if not key:
                continue  # a new row marked for deletion: nothing to do
            sheet.delete_keys.append(key)
        elif action:
            sheet.errors.append(f"Row {line}: unknown _action {cells[ACTION]!r} (use 'delete' or leave blank)")
        else:
            sheet.rows.append(row)
    return sheet


def changeset_from_edits(base: Rows, sheet: EditSheet, schema: Schema) -> tuple[Changeset, list[str]]:
    """Turn an edited sheet into a changeset relative to the version it was exported from."""
    errors = list(sheet.errors)
    base_ix = index(base, schema)
    for row in sheet.rows:
        k = row[schema.key]
        if k and k not in base_ix:
            errors.append(f"Record {k} does not exist. Leave the id blank for new records.")
    for k in sheet.delete_keys:
        if k not in base_ix:
            errors.append(f"Cannot delete record {k}: it does not exist.")
    if errors:
        return Changeset(), errors

    cs = diff(base, sheet.rows, schema, missing_means_delete=False)
    for k in sheet.delete_keys:
        cs.updates.pop(k, None)
        cs.deletes[k] = base_ix[k]
    return cs, []
